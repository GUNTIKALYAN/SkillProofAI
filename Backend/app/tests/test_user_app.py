import re
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.conftest import register
from app.main import app

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FRONTEND = Path(__file__).resolve().parents[3] / "Frontend"
_HEX = re.compile(r"#[0-9A-Fa-f]{3,8}")


def _resume_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Skills")
    document.add_paragraph("Python")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_history_lists_only_the_owner_and_saves_a_company_url(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.services.analysis_service.enqueue_analysis", lambda _analysis_id: None)
    with TestClient(app) as client:
        owner = {"Authorization": f"Bearer {register(client, 'owner@example.com')}"}
        other = {"Authorization": f"Bearer {register(client, 'reader@example.com')}"}
        created = client.post(
            "/api/analyses",
            headers=owner,
            data={"jd_text": "Requirements\nPython\nFastAPI\n", "github_username": "octocat"},
            files={"resume": ("resume.docx", _resume_bytes(), DOCX_TYPE)},
        )
        assert created.status_code == 202, created.text
        job_id = created.json()["job_id"]
        listing = client.get("/api/analyses", headers=owner)
        assert listing.status_code == 200
        assert [row["id"] for row in listing.json()] == [job_id]
        assert client.get("/api/analyses", headers=other).json() == []
        denied = client.patch(
            f"/api/analyses/{job_id}/company-url",
            headers=other,
            json={"name": "Acme", "url": "https://acme.example"},
        )
        assert denied.status_code == 404
        saved = client.patch(
            f"/api/analyses/{job_id}/company-url",
            headers=owner,
            json={"name": "Acme", "url": "https://acme.example"},
        )
        assert saved.status_code == 200, saved.text
        assert saved.json()["result"]["company_urls"] == [{"name": "Acme", "url": "https://acme.example"}]
        rejected = client.patch(
            f"/api/analyses/{job_id}/company-url",
            headers=owner,
            json={"name": "Acme", "url": "javascript:alert(1)"},
        )
        assert rejected.status_code == 422


def test_user_pages_are_static_and_do_not_inject_html() -> None:
    with TestClient(app) as client:
        for name in ("login.html", "signup.html", "app.html", "history.html", "chat.html", "compare.html", "admin.html"):
            page = client.get(f"/{name}")
            assert page.status_code == 200
            stem = name.removesuffix(".html")
            assert f"js/{stem}.js" in page.text
        assert client.get("/css/tokens.css").status_code == 200
        home = client.get("/", follow_redirects=False)
        assert home.status_code == 307
        assert home.headers["location"] == "/app.html"
    for path in (FRONTEND / "js").glob("*.js"):
        text = path.read_text(encoding="utf-8")
        assert "innerHTML" not in text
        assert "insertAdjacentHTML" not in text
    base = (FRONTEND / "css" / "base.css").read_text(encoding="utf-8")
    assert _HEX.search(base) is None
